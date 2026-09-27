import datetime as dt
import os
from decimal import Decimal

from openpyxl import load_workbook

from spendwise.db.models import (
    Account,
    Category,
    FamilyMember,
    ImportBatch,
    ImportBatchStatus,
    ReviewStatus,
    Transaction,
    TransactionDirection,
    TransactionType,
)
from spendwise.export.backup import create_backup, restore_backup
from spendwise.export.excel_export import _safe_cell, generate_excel_report
from spendwise.export.package import export_package, import_package, preview_package
from spendwise.analytics.metrics import compute_period_metrics


def _setup_member_with_txn(session, member_name="a", account_alias="acc1", fingerprint="fp1"):
    member = FamilyMember(name=member_name)
    session.add(member)
    session.flush()
    account = Account(family_member_id=member.id, alias=account_alias)
    session.add(account)
    session.flush()
    batch = ImportBatch(account_id=account.id, label="t", status=ImportBatchStatus.COMMITTED)
    session.add(batch)
    session.flush()
    category = session.query(Category).filter(Category.parent_id.isnot(None)).first()
    txn = Transaction(
        import_batch_id=batch.id, account_id=account.id, family_member_id=member.id,
        txn_date=dt.date(2026, 8, 1), amount=Decimal("50.00"), direction=TransactionDirection.DEBIT,
        original_currency="SAR", base_amount=Decimal("50.00"), txn_type=TransactionType.EXPENSE,
        category_id=category.id, merchant_sanitized="متجر", review_status=ReviewStatus.APPROVED,
        row_fingerprint=fingerprint,
    )
    session.add(txn)
    session.commit()
    return member, account


def test_safe_cell_neutralizes_formula_injection():
    assert _safe_cell("=SUM(A1:A10)").startswith("'")
    assert _safe_cell("normal text") == "normal text"
    assert _safe_cell("+1+1") .startswith("'")


def test_generate_excel_report_produces_valid_workbook(db_session, tmp_path):
    member, account = _setup_member_with_txn(db_session)
    metrics = compute_period_metrics(db_session, member.id, "2026-08")
    out_path = tmp_path / "report.xlsx"
    generate_excel_report(str(out_path), member.name, metrics)
    assert out_path.exists()
    wb = load_workbook(str(out_path))
    assert "الملخص" in wb.sheetnames


def test_export_package_contains_only_approved_sanitized_fields(db_session):
    member, account = _setup_member_with_txn(db_session)
    package = export_package(db_session, member.id)
    assert package["family_member_name"] == member.name
    assert len(package["transactions"]) == 1
    txn = package["transactions"][0]
    assert "merchant_sanitized" in txn
    assert "row_fingerprint" in txn
    assert "raw_description" not in txn
    assert "iban" not in str(txn).lower()


def test_export_package_includes_auto_approved_transactions(db_session):
    """Regression test: export_package must include AUTO_APPROVED rows too,
    not just APPROVED — otherwise high-confidence transactions silently
    disappear from exported packages and backups."""
    member, account = _setup_member_with_txn(db_session, fingerprint="fp-auto")
    txn = db_session.query(Transaction).filter_by(row_fingerprint="fp-auto").one()
    txn.review_status = ReviewStatus.AUTO_APPROVED
    db_session.commit()

    package = export_package(db_session, member.id)
    assert len(package["transactions"]) == 1


def test_package_import_skips_duplicates_and_scopes_by_family_member(db_session):
    """Regression test for the cross-family-member false-positive bug: a
    package should only be treated as 'already imported' relative to the
    TARGET family member, not globally."""
    member_a, account_a = _setup_member_with_txn(db_session, member_name="a", account_alias="shared_alias", fingerprint="shared-fp")
    package = export_package(db_session, member_a.id)

    # Importing into a brand-new member with an account alias that happens to
    # match member_a's alias must NOT be treated as already-imported.
    result = import_package(db_session, package, "b")
    db_session.commit()
    assert result["imported_count"] == 1
    assert result["skipped_count"] == 0

    # Re-importing into the SAME new member should now correctly skip it as
    # a true duplicate.
    result2 = import_package(db_session, package, "b")
    db_session.commit()
    assert result2["imported_count"] == 0
    assert result2["skipped_count"] == 1


def test_preview_package_reports_counts(db_session):
    member, account = _setup_member_with_txn(db_session)
    package = export_package(db_session, member.id)
    preview = preview_package(db_session, package, target_family_member_id=None)
    assert preview["total_transactions"] == 1
    assert preview["new_count"] == 1


def test_backup_and_restore_roundtrip(tmp_path):
    db_path = tmp_path / "spendwise.db"
    settings_path = tmp_path / "app_settings.json"
    db_path.write_text("fake db content")
    settings_path.write_text('{"base_currency": "SAR"}')

    backup_path = tmp_path / "backup.zip"
    create_backup(str(backup_path), str(db_path), str(settings_path))
    assert backup_path.exists()

    db_path.write_text("corrupted")
    restore_backup(str(backup_path), str(db_path), str(settings_path))
    assert db_path.read_text() == "fake db content"
    assert os.path.exists(str(db_path) + ".pre_restore_backup.db")
