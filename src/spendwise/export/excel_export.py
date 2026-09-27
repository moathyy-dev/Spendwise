"""
Generates RTL Excel reports/exports with formula-injection defense: any
cell value that starts with a formula-trigger character gets a leading
apostrophe so Excel/LibreOffice treat it as plain text, never as a formula.
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

_FORMULA_TRIGGER_CHARS = ("=", "+", "-", "@")


def _safe_cell(value):
    if isinstance(value, str) and value.startswith(_FORMULA_TRIGGER_CHARS):
        return "'" + value
    return value


def _style_sheet_rtl(ws):
    ws.sheet_view.rightToLeft = True


def _header_row(ws, headers: list[str], row: int = 1):
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col_idx, value=_safe_cell(header))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="right")


def generate_excel_report(output_path: str, family_member_name: str, metrics) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "الملخص"
    _style_sheet_rtl(ws)

    ws.append([_safe_cell(f"تقرير {family_member_name} — {metrics.month}")])
    ws.append([])
    _header_row(ws, ["المؤشر", "القيمة"], row=3)
    rows = [
        ("إجمالي المصروفات", float(metrics.total_expenses)),
        ("إجمالي الدخل", float(metrics.total_income)),
        ("صافي الادخار", float(metrics.net_savings)),
        ("إجمالي التحويلات", float(metrics.transfers_total)),
    ]
    for i, (label, value) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=_safe_cell(label)).alignment = Alignment(horizontal="right")
        ws.cell(row=i, column=2, value=value)

    ws2 = wb.create_sheet("حسب التصنيف")
    _style_sheet_rtl(ws2)
    _header_row(ws2, ["التصنيف", "المبلغ"])
    for i, (cat_id, amount) in enumerate(sorted(metrics.category_totals.items(), key=lambda kv: kv[1], reverse=True), start=2):
        name = metrics.category_names.get(cat_id, "غير مصنّف")
        ws2.cell(row=i, column=1, value=_safe_cell(name)).alignment = Alignment(horizontal="right")
        ws2.cell(row=i, column=2, value=float(amount))

    wb.save(output_path)
    return output_path


def generate_transactions_excel(output_path: str, transactions: list) -> str:
    """`transactions` items expose .txn_date, .amount, .direction,
    .merchant_sanitized, .category (with name_ar), .txn_type — all
    sanitized, non-sensitive fields only."""
    wb = Workbook()
    ws = wb.active
    ws.title = "المعاملات"
    _style_sheet_rtl(ws)
    _header_row(ws, ["التاريخ", "المبلغ", "الاتجاه", "التاجر", "التصنيف", "النوع"])

    for i, t in enumerate(transactions, start=2):
        category_name = t.category.name_ar if getattr(t, "category", None) else "غير مصنّف"
        direction_ar = "مدين" if getattr(t.direction, "value", t.direction) == "debit" else "دائن"
        ws.cell(row=i, column=1, value=t.txn_date.isoformat())
        ws.cell(row=i, column=2, value=float(t.amount))
        ws.cell(row=i, column=3, value=_safe_cell(direction_ar)).alignment = Alignment(horizontal="right")
        ws.cell(row=i, column=4, value=_safe_cell(t.merchant_sanitized or "")).alignment = Alignment(horizontal="right")
        ws.cell(row=i, column=5, value=_safe_cell(category_name)).alignment = Alignment(horizontal="right")
        ws.cell(row=i, column=6, value=_safe_cell(getattr(t.txn_type, "value", t.txn_type))).alignment = Alignment(horizontal="right")

    wb.save(output_path)
    return output_path
